"""Build the argument brief and the committed, source-led book outline."""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "book"
OUT = BOOK / "drafts"

BRIEF = [
    [
        ("", "A song called Lockdown finished first in one MusicLab world and fortieth in another. The recording was the same, competing against the same songs. What changed was the sequence of listeners and the popularity information they encountered. Salganik, Dodds, and Watts created the parallel markets; the rankings are also described in Salganik's own account. Shadow Futures begins with that documented result and asks what we can learn about success when ordinary economic life gives us only one history (Salganik, Dodds, and Watts, 2006; Salganik, Bit by Bit)."),
        ("The comparison budget", "Imagine two software suppliers, Aster and Birch. They are explicitly fictional. Their work and product characteristics are observable, but the effect of those inputs on winning contracts is unknown. An early contract can build references and access that influence later contracts. For each award, take the chance that it goes to someone other than the current favorite. Add those chances over the sequence. I call this total the comparison budget. The favorite can change. The budget adds opportunities outside the favorite; it does not count actual rival victories or independent experiments."),
        ("", "In my specified allocation model, that budget places an upper bound on the information the award record supplies about productive input effects. A large budget still requires informative differences between alternatives. A small chance repeated indefinitely can accumulate without limit. Concentration or a low current chance alone therefore does not establish the strongest result. The relevant question is how much comparison the whole allocation process preserves (Erlic, v5, Sections 2 and 3)."),
        ("Shadow futures", "A shadow future is another feasible allocation history with the same starting conditions, rules, and underlying effects of productive inputs. In one, Birch receives the early contract and develops differently. Later effort can respond to earlier outcomes under the same response rule. These replays reveal historical variation. A separate intervention, such as changing effort, defines the contribution we want to estimate. An imagined better economy is not evidence merely because we can describe it (Erlic, v5, Section 4)."),
        ("", "The strongest result requires common rules, the same possible next winners under each explanation, controlled differences between their predictions, and a finite lifetime comparison budget. Under those conditions, there is no single method whose contribution estimates become reliably accurate under every candidate explanation as one history grows. I call the unresolved attribution contribution uncertainty. Work can still matter. Extra outcome measurements carrying information about its effect must be counted separately. Forecasts can nevertheless converge, using Blackwell and Dubins's classical result (1962), while the attribution limit remains (Erlic, v5, Section 3 and Appendix B)."),
    ],
    [
        ("What is new", "I propose an economic information bound connecting the comparison budget to what reward histories reveal about the effects of work, quality, and investment. It limits precision in a finite record and, under further conditions, reliable attribution as history extends indefinitely. Alchian (1950) already explained profit-based selection under uncertainty. Increasing returns and cumulative advantage are established ideas (Arthur, 1989; Merton, 1968). Estimation failure in a reinforced urn predates this paper (Le Goff and Soulier, 2017), as does accurate prediction alongside inconsistent parameter estimation (Zhang, 2004). The contribution is this defined object, bound, and economic interpretation. Its probability foundations are classical; priority requires wider review."),
        ("AI and work", "AI makes the distinction between activity and comparison consequential. Frontier training and deployment depend on compute, infrastructure, and distribution. Inference means running a trained model to produce an output. A growing volume of outputs can reveal performance on encountered tasks without independently testing alternative suppliers, training strategies, or economic arrangements. AI can also lower the cost of those tests. Whether it expands the comparison budget is an empirical and institutional question. The work chapter pairs observed productivity gains in customer support with METR's account of developers avoiding tasks that might be assigned to a no-AI condition (Brynjolfsson, Li, and Raymond, 2025; METR, 2026)."),
        ("Antitrust", "An access rule or merger can change which alternatives receive interpretable trials. This supplies a possible informational rationale for intervention alongside established price, quality, and innovation concerns. I use epistemic monopoly for control over relevant comparison production, explicitly crediting Meghani's prior use of the term (2021). An intervention must improve a consequential decision enough to justify its costs. Internal experiments may preserve learning, and independent routes may be expensive. This economic argument does not itself establish legal liability or remove existing statutory requirements (Erlic, v5, Sections 4 and 5)."),
        ("Taxation", "There are two arguments. First, economies that generate indistinguishable observable histories can attribute different contributions to the same reward. A rule using those histories cannot divide every reward exactly into earned income and positional advantage in both economies. This separate hidden-position argument does not invalidate optimal taxation (Erlic, v5, Section 5 and Appendix E; Piketty, Saez, and Stantcheva, 2014; Rothschild and Scheuer, 2016). Second, conduct removing informative trials may impose an external cost. A corrective tax needs evidence of that causal link, a valuable affected decision, private incentives that miss part of the social cost, and an instrument worth its cost. Arrow helps supply the additional welfare foundation (1962, 1963)."),
        ("The book", "The narrative returns to Lockdown and to Aster and Birch as the stakes move from cultural success to AI procurement, work, and policy. Published experiments, court records, and economic history carry the real cases; interviews are not required. The closing question is how institutions can preserve enough informative opportunity to discover better arrangements while capturing the benefits of scale. The technical source is my SSRN preprint and its October 2026 v5 manuscript, not an accepted journal publication."),
    ],
]

CHAPTERS = [
    {
        "number": 1, "title": "The winner and the explanation",
        "paragraphs": [
            "Open with the two rankings of Lockdown, already established in the prologue. Examine the explanation a reader might give after seeing only the winning world. The same song's contrasting outcomes make the difference between appeal and realized success concrete without claiming that quality is irrelevant.",
            "Move to van de Rijt and colleagues' Kickstarter experiment: researchers selected 200 new unfunded projects and donated to 100 chosen at random. Follow its comparison of subsequent support. Retain the paper's finding of diminishing returns to larger initial advantages, which limits a simple runaway-success story. Introduce Merton's Matthew effect and Akerlof's hidden-quality problem as distinct predecessors. Ask what would distinguish productive inputs from the effect of having already succeeded.",
            "Aster and Birch first appear in a short, labeled thought experiment. Both perform real work. One early contract changes references and future access. No invented dialogue or personal history is attached to them. The chapter closes before assigning either supplier an exact share of merit."
        ],
        "sources": "V5 citations: Salganik, Dodds, and Watts (2006); van de Rijt et al. (2014); Merton (1968); Akerlof (1970). Added source: Salganik's first-person MusicLab account.",
    },
    {
        "number": 2, "title": "The comparison budget",
        "paragraphs": [
            "Stay with the fictional procurement desk choosing Aster or Birch. Follow a small sequence of awards and state the model-assumed chance of an award outside the current favorite in ordinary words. Add those chances conceptually, with no formula. Distinguish the modeled chances from market shares, realized wins, and a count of independent trials.",
            "Explain why mostly repeating the favored choice may supply little information about input effects, even while the service improves and transactions multiply. Define the particular information question before describing its scarcity. A high comparison budget is an opportunity for information, not a guarantee of identification.",
            "Give two endings. In one, a persistent chance remains and comparison continues accumulating. In the other, chances diminish fast enough for their lifetime total to be finite. Describe strong reinforcement as one sufficient setting, crediting the earlier reinforced-process literature. Close with the cost of allocating some contracts to less certain suppliers."
        ],
        "sources": "V5 citations: Bar-Yam (2016); Pemantle (2007); Oliveira (2009); Gottfried and Grosskinsky (2024); Kasy and Sautmann (2021). V5 Sections 2 and 3 and Appendix D.",
    },
    {
        "number": 3, "title": "Shadow futures",
        "paragraphs": [
            "Replay the procurement example with the same underlying parameters, starting state, and response rules, but different early awards. Let later behavior change only as the unchanged rules prescribe. Separate these histories from an intervention that changes an input or a policy. Return briefly to MusicLab as an observed analogue of parallel histories, not a test of the theorem.",
            "Explain the finite-budget result in prose and preserve its conditions. Candidate explanations must use the common recorded experiment, retain the same possible next recipients, and differ in a way controlled by remaining comparison. Additional informative outcomes belong in the experiment. The conclusion is a failure of universal consistent recovery, not identical probabilities, total ignorance, or a guarantee that every observed market is permanently unknowable.",
            "Use Blackwell and Dubins to explain how forecasts can converge while attribution remains unresolved in v5's setting. Credit Le Goff and Soulier and Zhang for prior estimation and prediction results. A single grouped technical endnote credits Kakutani, Kabanov and colleagues, Jacod and Shiryaev, Gabriyelyan, and Le Cam and Yang. Close with what an independent, informative replication could add."
        ],
        "sources": "V5 citations: Blackwell and Dubins (1962); Le Goff and Soulier (2017); Zhang (2004); Kakutani (1948); Kabanov, Liptser, and Shiryaev (1977); Jacod and Shiryaev (2003); Gabriyelyan (2011); Le Cam and Yang (2000). V5 Appendices A to E.",
    },
    {
        "number": 4, "title": "The histories that became infrastructure",
        "paragraphs": [
            "Commit to Bell's 1956 patent licensing decree as the historical case. Start with what the decree required and follow the published study of subsequent invention. Explain the study's finding that innovation rose outside telecommunications, together with its different result inside that sector. Do not collapse the 1956 decree into the later Bell breakup.",
            "Connect this documented access change to Arthur's account of increasing returns and historical selection. Put the argument in dialogue with Alchian's 1950 account of profit-based selection under uncertainty: selection and recovery of productive contribution are different questions. Credit David's QWERTY paper in a short note acknowledging that the keyboard's empirical interpretation is disputed. Merton supplies another established account of cumulative advantage.",
            "The interpretive question is whether institutions permit alternatives to develop usable records. Bell's innovation findings do not estimate its comparison budget or establish a finite lifetime budget. Bring back Birch as a supplier gaining a route to demonstrate its capabilities. Close by distinguishing access to existing knowledge from opportunities to produce new comparative evidence."
        ],
        "sources": "V5 citations: Arthur (1989); David (1985); Merton (1968); van de Rijt et al. (2014); Hayek (2002). Added sources: Alchian (1950); Watzinger, Fackler, Nagler, and Schnitzer (2020).",
    },
    {
        "number": 5, "title": "Who gets to be tested",
        "paragraphs": [
            "Open with the publicly documented Google search distribution dispute and the September 2025 remedies involving certain data and syndication access. Separate allegations, findings, remedies, and the book's economic interpretation. Use the dated court record; check subsequent litigation when the chapter is drafted.",
            "Credit Zahra Meghani before introducing the narrower use of epistemic monopoly: control over comparison production relative to a defined contribution question and information set. Her knowledge-governance argument is an antecedent. The book does not claim to have coined the term. Contrast control of relevant trials with possession of a large dataset.",
            "Return to Aster and Birch on several storefronts sharing one selection system. Then consider a single operator that maintains genuinely informative pilots. These thought experiments explain why organization count and comparison opportunities differ. Close by asking who can check whether an incumbent's private experiments answer the public's consequential questions."
        ],
        "sources": "V5 citations: Meghani (2021); Hayek (2002). V5 Sections 4 and 5. Added sources: Google's search case documents and DOJ's September 2, 2025 remedies summary.",
    },
    {
        "number": 6, "title": "Who deserves the winnings",
        "paragraphs": [
            "Return to the composer of Lockdown only through the documented recording and experimental rankings. Do not invent the musician's reaction. Ask what a reward establishes about effort, appeal, luck, and position when we have seen only one resulting market.",
            "Separate causal contribution, contractual entitlement, incentive design, and moral desert. Use entrepreneurship as experimentation to explain why large prospective prizes can be useful. Explain how established tax models address effort, bargaining, and rent-seeking without requiring an exact causal audit of every person's income.",
            "Apply the distinction to Aster's fictional owners and Birch's workers. Both can have acted productively even when their rewards diverge. The book argues for care in retrospective attribution and then states its normative choices openly. It does not infer an optimal tax rate, equal merit, or an exact luck percentage. Close by moving from the moral defense of winners to the physical investment needed for AI."
        ],
        "sources": "V5 citations: Kerr, Nanda, and Rhodes-Kropf (2014); Piketty, Saez, and Stantcheva (2014); Rothschild and Scheuer (2016); Merton (1968); Salganik, Dodds, and Watts (2006).",
    },
    {
        "number": 7, "title": "The intelligence factory",
        "paragraphs": [
            "Build the opening from the IEA's documented data-center infrastructure account, without claiming to have visited a facility. Follow the physical requirements of a computing cluster: accelerators, network capacity, power, cooling, and capital. Use Stanford's 2026 AI Index for the industry context. Date all investment and model-capability claims.",
            "Explain frontier training, post-training, and inference as different economic activities. Frontier describes the leading capability at a given date, not an enduring group of firms. Scale can generate useful capability and lower costs while raising the expense of independent entry. Keep that tension rather than asserting that a large training run proves monopoly or information destruction.",
            "Aster and Birch now purchase AI capacity under explicitly specified hypothetical arrangements. Compare access to a common infrastructure service with an exclusive route to users. Arrow supplies the established appropriation problem, and entrepreneurship research supplies the incentive argument. Close by following a trained system from the factory into repeated use."
        ],
        "sources": "V5 citation: Kerr, Nanda, and Rhodes-Kropf (2014). Added sources: Arrow (1962); Stanford AI Index (2026); IEA Key Questions on Energy and AI (2026).",
    },
    {
        "number": 8, "title": "A billion answers",
        "paragraphs": [
            "Follow an explicitly hypothetical inference request through the supplier choice, generated answer, user response, and next routing decision. Repetition can reveal useful performance on encountered tasks. It need not identify the contribution of inputs or the outcome of an alternative training and distribution history.",
            "Distinguish evaluating fixed models on comparable tasks from learning the returns to different economic institutions. Shared model ancestry and a shared ranking can make seemingly separate services dependent. Conversely, routing experiments and independent deployment can add informative evidence. Logged routing probabilities and informative reward labels must be modeled; the allocation-only theorem cannot simply be applied after those observations are ignored.",
            "Return to the prediction and explanation distinction, crediting Blackwell and Dubins and Zhang. Give AI's positive case equal analytical weight: it may make independent tests cheaper. Simulated worlds can test implications of assumptions; their realism requires external validation. Close with the people whose work supplies training examples and whose jobs change when deployment expands."
        ],
        "sources": "V5 citations: Bar-Yam (2016); Blackwell and Dubins (1962); Zhang (2004). V5 Section 3. Added context: Stanford AI Index (2026).",
    },
    {
        "number": 9, "title": "The first rung",
        "paragraphs": [
            "Commit to the published customer-support deployment studied by Brynjolfsson, Li, and Raymond. Explain its average productivity gain and the different effects across workers, including evidence of learning. Follow the observed workplace design; invent no worker dialogue. Distinguish task output from wages, headcount, and an economy-wide employment effect.",
            "Then turn to METR's early-2025 developer experiment and its February 2026 account of changing selection into the follow-up study. Developers sometimes avoided participation or tasks that might require working without AI; lower pay and time measurement also affected interpretation. The lesson is that adoption can change which comparisons are available. It is not a demonstration that v5's finite-budget conditions hold.",
            "Use Birch's hypothetical junior staff to investigate a future question: can automating initial tasks remove routes for developing and demonstrating expertise? The customer-support evidence also supports the possibility of faster learning. Close with a concrete institutional choice between monitored augmentation, immediate replacement, and preserving a limited comparison group, with their costs stated."
        ],
        "sources": "Added sources: Brynjolfsson, Li, and Raymond (2025, published version); METR early-2025 trial; METR update of February 24, 2026. V5 supplies the question, not an empirical finding about labor markets.",
    },
    {
        "number": 10, "title": "Competition for future evidence",
        "paragraphs": [
            "Revisit the Google search remedies and Bell's licensing experience. State that competition policy already considers quality, innovation, entry, and labor conditions. The additional proposed rationale concerns the production of evidence useful for a specified decision. Economic desirability and legal authority remain separate questions.",
            "Commit to a fictional public buyer choosing between Aster and Birch for an AI-assisted service. Specify eligible suppliers, comparable tasks, outcome definitions, retention of trial records, and a purchasing decision. Compare an audit, portability, a limited pilot, and independent procurement routes. Blackwell helps explain decision value; Kasy and Sautmann explain why adaptive allocation involves a learning choice.",
            "Judge remedies by attainable decision improvement, implementation costs, scale benefits, and uncertainty about how to interpret the evidence. Several trials sharing one uncalibrated measure can fail together. Credit Meghani and Hayek for antecedents. The chapter ends with a conditional institutional recommendation rather than a legal finding about a named firm."
        ],
        "sources": "V5 citations: Meghani (2021); Hayek (2002); Kasy and Sautmann (2021). Added foundation: Blackwell (1953). Policy extension: Erlic's unpublished Competition Policy for Informative Market Access draft. Added legal source: DOJ 2023 Merger Guidelines, Guideline 2.",
    },
    {
        "number": 11, "title": "Taxing lost comparison",
        "paragraphs": [
            "Begin with the fictional buyer's expense of replacing a lost independent supplier trial. Introduce no real-company tax estimate. State the separate v5 argument: when hidden position can offset a changed input effect, identical observable laws can assign different contributions to a reward. A rule using that record cannot implement an exact earned-versus-positional decomposition in every such economy.",
            "Connect this limit to established public-finance work without claiming it refutes optimal taxation. State the chosen distributional principles separately. Then develop the additional corrective-tax argument from a defined action to lost evidence, a worse consequential decision, and an external cost that private incentives fail to internalize.",
            "Compare an access remedy, a testing payment, and funded evaluation. Retain the companion draft's insights that public trials can crowd out private testing and that feasible replacement can limit the cost of lost comparison. Compute or profit alone is an unreliable proxy for this harm. Close by explaining which measurements would be necessary before choosing a real instrument or rate."
        ],
        "sources": "V5 citations: Piketty, Saez, and Stantcheva (2014); Rothschild and Scheuer (2016); Kerr, Nanda, and Rhodes-Kropf (2014). Added foundation: Arrow (1962). Extension: Erlic's unpublished Taxing Lost Comparison draft, whose additional model assumptions must be stated.",
    },
    {
        "number": 12, "title": "An economics of possible worlds",
        "paragraphs": [
            "Open at Rothamsted's documented long-term agricultural experiments. Explain the institutional commitment needed to preserve useful comparisons while adapting experiments to changing farming conditions. This is a constructive historical example of producing evidence, not an application of v5's theorem to wheat.",
            "Return to Aster and Birch and choose a costed, time-limited evaluation design rather than indefinite exploration. Explain how experimental access, outcome definitions, compatible records, and independent evaluation can work together. Recognize the established economics of learning, counterfactual analysis, discovery, and experimental design.",
            "End with Lockdown. In the experimental setting its alternative outcomes were observed; in ordinary markets many alternatives remain unrealized. The proposal is to measure and preserve valuable comparison opportunities where feasible. The reader should leave able to distinguish an observed winner, an explanation of contribution, and a justified institutional choice. The epilogue returns to that distinction in one final procurement decision."
        ],
        "sources": "V5 citations: Hayek (2002); Kasy and Sautmann (2021); Blackwell and Dubins (1962); Salganik, Dodds, and Watts (2006). Added sources: Rothamsted's history and account of adapting its experiments; Blackwell (1953).",
    },
]


def setup(title: str, brief: bool = False) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.72 if brief else 0.78)
    sec.left_margin = sec.right_margin = Inches(0.82)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11.5 if brief else 11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.07 if brief else 1.08
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 19), ("Subtitle", 10.5), ("Heading 1", 14), ("Heading 2", 11.5)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(9 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(5)
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
    for style in [normal, doc.styles["Title"], doc.styles["Subtitle"], doc.styles["Heading 1"], doc.styles["Heading 2"]]:
        ppr = style.element.get_or_add_pPr()
        for border in list(ppr.findall(qn("w:pBdr"))):
            ppr.remove(border)
    doc.core_properties.title = title
    doc.core_properties.author = "Martin Erlic"
    doc.add_paragraph(title, "Title")
    doc.add_paragraph("Martin Erlic   16 October 2026", "Subtitle")
    return doc


def paragraph(doc: Document, text: str, small: bool = False) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.widow_control = True
    if small:
        p.paragraph_format.space_after = Pt(5)
        for run in p.runs:
            run.font.size = Pt(9.5)


def pagebreak(doc: Document) -> None:
    doc.add_page_break()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    brief = setup("Shadow Futures Argument Brief", brief=True)
    for number, content in enumerate(BRIEF):
        if number:
            pagebreak(brief)
        for heading, text in content:
            if heading:
                brief.add_paragraph(heading, "Heading 2")
            paragraph(brief, text)
    brief_path = OUT / "Shadow_Futures_Argument_Brief.docx"
    brief.save(brief_path)

    outline = setup("Shadow Futures Book Outline")
    paragraph(outline, "An 80000 word popular economics book about what reinforced markets can reveal about contribution and how institutions can preserve useful comparison. The argument moves from one song's contrasting outcomes to the choices governing AI, work, competition, and taxation. The book uses published research and public records. Its first draft requires no interviews.")
    outline.add_paragraph("The narrative throughline", "Heading 1")
    paragraph(outline, "The recurring documented object is Lockdown by 52Metro. Its first-place and fortieth-place outcomes in MusicLab frame the prologue and return in Chapters 1, 3, 6, and 12. Salganik's own account supplies the specific rankings. The story follows what the observer can infer in each setting; the book invents no experiences or reactions for the musicians.")
    paragraph(outline, "Aster and Birch are recurring fictional software suppliers. Their first contract introduces feedback; procurement reveals the comparison budget; AI deployment raises the stakes; the last chapters compare institutional remedies. Every appearance is labeled as a thought experiment. Changing the early awards while holding economic rules fixed is a replay. Changing an input or policy is a separate intervention.")
    paragraph(outline, "The real-case sequence is fixed: MusicLab, the success-breeds-success field experiments, Bell's 1956 licensing decree, Google search distribution and remedies, AI infrastructure, customer-support deployment, METR's developer studies, and Rothamsted. Each case supplies evidence for a stated mechanism or institutional choice. None is presented as an empirical proof of an infinite-horizon theorem.")
    outline.add_paragraph("Opening and word budget", "Heading 1")
    paragraph(outline, "Prologue 2500 words: two published outcomes of the same recording, followed by the missing comparison in ordinary markets. Introduction 3500 words: the attribution question and the three defined concepts. Chapters 1 to 12 total 72000 words. Epilogue 2000 words: return to the song and the final procurement decision. Notes and references are additional.")
    paragraph(outline, "The four movements are learning the concepts in Chapters 1 to 3; history, power, and desert in Chapters 4 to 6; AI and work in Chapters 7 to 9; and institutions in Chapters 10 to 12. The author narrates documented choices and explains mechanisms. Sensory scenes, dialogue, motives, and private experiences require an identified record.")

    for index in range(0, len(CHAPTERS), 2):
        pagebreak(outline)
        for chapter in CHAPTERS[index:index + 2]:
            outline.add_paragraph(f"{chapter['number']} {chapter['title']}", "Heading 1")
            for text in chapter["paragraphs"]:
                paragraph(outline, text)
            paragraph(outline, chapter["sources"], small=True)

    pagebreak(outline)
    outline.add_paragraph("The earlier work on information and uncertainty", "Heading 1")
    lineage = [
        ("Alchian in 1950", "Uncertainty, Evolution, and Economic Theory is the closest title match to the remembered 1950s paper. Alchian explains how realized profits can select behavior even without individually optimal foresight. Shadow Futures asks what a resulting history can reveal about productive contributions when past rewards alter later opportunities. This is a further attribution question, not a general refutation of market selection. Alchian is an added reference, absent from v5's 24-item bibliography."),
        ("Blackwell in 1953", "Equivalent Comparisons of Experiments gives a decision-based way to compare experiments. Roughly, one experiment is at least as informative as another when it permits at least as good a decision across the theorem's decision problems, with a corresponding statistical transformation characterization. This is a foundation for valuing preserved comparison, not a theorem that more observations always justify their cost. It is an addition to the book's bibliography, not one of v5's 24 references."),
        ("Blackwell and Dubins in 1962", "Merging of Opinions with Increasing Information is the direct classical source cited in v5. Under its absolute-continuity condition, conditional predictions merge as information accumulates. It does not say everyone converges to truth or that causal effects become identifiable. V5 combines this result with its particular single-history allocation limit."),
        ("Arrow in 1962", "Economic Welfare and the Allocation of Resources for Invention studies knowledge production under uncertainty and problems of appropriating its benefits. It supports the welfare discussion about private incentives and useful information. It does not establish the comparison budget or price its depletion. This is an added book reference."),
        ("Arrow in 1963", "Uncertainty and the Welfare Economics of Medical Care examines how uncertainty about illness and treatment, information differences, and risk-bearing shape medical institutions. This is another uncertainty classic from the period. It is a comparison for the institutional argument, not evidence that healthcare satisfies v5's theorem. It is an added book reference."),
        ("The wider lineage", "Knight's Risk, Uncertainty and Profit is from 1921, outside the 1950s and 1960s. Hayek's competition-as-discovery lecture dates to 1968; v5 cites the 2002 translation. Merton's Matthew effect is from 1968. Kakutani's mathematical predecessor is from 1948. Akerlof's lemons paper is from 1970. These are different contributions and should not be compressed into a claim that uncertainty or market learning was previously overlooked."),
    ]
    for heading, text in lineage:
        outline.add_paragraph(heading, "Heading 2")
        paragraph(outline, text)

    pagebreak(outline)
    outline.add_paragraph("The citation standard", "Heading 1")
    paragraph(outline, "All 24 references in the canonical October 2026 v5 are retained below and assigned to chapters in research/v5_reference_register.json. Relevant antecedents receive author-date attribution in the explanatory text or a specific endnote. Grouped notes credit mathematical sources without introducing equations. A planned location is not a completed verification of a source or a drafted chapter.")
    paragraph(outline, "Before a chapter's factual or novelty claim is finalized, check the relevant full source and add its page or section to the working citation record. Identify unpublished Erlic extensions as unpublished drafts and state their additional assumptions. Preserve the distinction between a theoretical result, an empirical finding, an illustrative analogy, and a policy proposal. Every v5 source has a substantive role or technical-lineage note; none is silently omitted.")
    outline.add_paragraph("The v5 references", "Heading 1")
    references = json.loads((BOOK / "research" / "v5_reference_register.json").read_text(encoding="utf-8"))["references"]
    for index, item in enumerate(references):
        if index == 12:
            pagebreak(outline)
            outline.add_paragraph("The v5 references continued", "Heading 1")
        paragraph(outline, item["citation"] + " Book chapters " + ", ".join(map(str, item["planned_chapters"])) + ".", small=True)

    pagebreak(outline)
    outline.add_paragraph("Additional primary sources for the book", "Heading 1")
    additions = [
        ("Alchian (1950)", "Uncertainty, Evolution, and Economic Theory. Journal of Political Economy 58(3), 211 to 221. https://doi.org/10.1086/256940. An original scan is retained in book/sources/classics."),
        ("Blackwell (1953)", "Equivalent Comparisons of Experiments. The Annals of Mathematical Statistics 24(2), 265 to 272. https://doi.org/10.1214/aoms/1177729032"),
        ("Arrow (1962)", "Economic Welfare and the Allocation of Resources for Invention. In The Rate and Direction of Inventive Activity, 609 to 626. NBER and Princeton University Press. https://www.nber.org/system/files/chapters/c2144/c2144.pdf"),
        ("Arrow (1963)", "Uncertainty and the Welfare Economics of Medical Care. American Economic Review 53(5), 941 to 973. https://assets.aeaweb.org/asset-server/files/9442.pdf"),
        ("Salganik", "Bit by Bit, MusicLab account in the open review edition, section 4.6.1. The first-person source for Lockdown's two ranks. https://www.bitbybitbook.com/en/running-experiments/exp-advice/zero-variable-cost/"),
        ("Watzinger, Fackler, Nagler, and Schnitzer (2020)", "How Antitrust Enforcement Can Spur Innovation: Bell Labs and the 1956 Consent Decree. American Economic Journal: Economic Policy 12(4), 328 to 359. https://doi.org/10.1257/pol.20190086"),
        ("Brynjolfsson, Li, and Raymond (2025)", "Generative AI at Work. Quarterly Journal of Economics 140(2), 889 to 942. Use this published version for the 15 percent finding. https://doi.org/10.1093/qje/qjae044"),
        ("METR (2025 and 2026)", "The early-2025 experienced-developer trial and We are Changing our Developer Productivity Experiment Design, February 24, 2026. The later post identifies selection and measurement problems. https://metr.org/blog/2026-02-24-uplift-update/"),
        ("US Department of Justice", "Google search case and September 2, 2025 remedies announcement. Separately, 2023 Merger Guidelines, Guideline 2, covers nonprice competition. https://www.justice.gov/opa/pr/department-justice-wins-significant-remedies-against-google"),
        ("Stanford HAI and IEA (2026)", "AI Index Report, Research and Development and Economy chapters; Key Questions on Energy and AI. https://hai.stanford.edu/ai-index/2026-ai-index-report and https://www.iea.org/reports/key-questions-on-energy-and-ai"),
        ("Rothamsted Research (2024)", "Why do we make changes to the long-term experiments at Rothamsted. Institutional account of the experiments begun in 1843 and their adaptation. https://www.rothamsted.ac.uk/news/why-do-we-make-changes-long-term-experiments-rothamsted"),
    ]
    for label, text in additions:
        paragraph(outline, label + ". " + text, small=True)
    outline_path = OUT / "Shadow_Futures_Committed_Book_Outline.docx"
    outline.save(outline_path)

    # Preserve editable source content and a durable agreement for future drafting.
    (BOOK / "research" / "book_plan_content.json").write_text(json.dumps({"brief": BRIEF, "chapters": CHAPTERS, "lineage": lineage, "added_sources": additions}, ensure_ascii=False, indent=2), encoding="utf-8")
    guidelines = """SHADOW FUTURES BOOK DRAFTING AGREEMENT

Author: Martin Erlic. Main-text target: 80000 words. No equations in the trade book.
Use the canonical October 2026 v5 identified in research/v5_source_provenance.json.
Both project copies match C:\\Shadow Futures\\Final byte for byte. Do not replace
the source manuscript or silently substitute an earlier SSRN version.

No interviews are required for the first draft. Use documented experiments,
published histories, original papers, court records, and institutional records.
The recurring real object is Lockdown by 52Metro in MusicLab. Aster and Birch
are explicitly fictional suppliers used only as labeled thought experiments.
Invent no dialogue, private motivations, visits, or personal reactions.

Keep all 24 v5 references in the live register. Credit relevant predecessors
where their ideas enter. Reserve mathematical lineage for explanatory notes.
Check the full source and record a page or section before finalizing a claim.
The inventory and chapter mapping are not a completed full-text literature audit.

Claim novelty for the specified comparison-budget information bound, its
finite-record precision limit, and its economic attribution interpretation.
Credit prior path dependence, urn estimation failure, predictive merging,
prediction with inconsistent parameter estimation, and epistemic monopoly.
Do not claim scholarly priority or proof certification without independent review.

Define the budget as cumulative modeled allocation chance outside the current
favorite. Distinguish it from actual rival wins, independent trial counts,
market shares, firm counts, and welfare. Infinite budget need not identify.
Small or shrinking chance alone does not establish a finite lifetime budget.
Count all informative observed outcomes in any application of the theorem.

Shadow futures preserve the parameter, initial state, and input/state response
rules; adaptive realized inputs can vary with history. Input and policy changes
are distinct interventions. Separate the finite-budget theorem from the
hidden-position observational-equivalence argument used in the tax discussion.

AI, corrective taxation, and remedy design are applications or extensions.
State their extra causal, empirical, welfare, and legal premises. Assess the
benefits of scale, private experiments, costs, and replacement testing.
Do not infer an actual firm's liability, optimal tax rate, or infinite-horizon
condition from a market-share chart, transaction count, or illustrative simulation.
"""
    (BOOK / "README.txt").write_text(guidelines, encoding="utf-8")
    summary = {"brief_words": sum(len(re.findall(r"\b[\w'-]+\b", t)) for page in BRIEF for _, t in page), "chapter_count": len(CHAPTERS), "v5_references_retained": len(references), "outputs": [str(brief_path), str(outline_path)]}
    (BOOK / "research" / "build_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary), flush=True)


if __name__ == "__main__":
    main()
